import os
import logging
import pypdf
from docx import Document
import tempfile
import asyncio
import base64
from config import Config
from doc_processor.nanonetocr import NanoNetsOCRProcessor
from rag.nanonet_ocr import ocr_list_pages
from rag.utils import crop_all_pages, encode_image
from PIL import Image

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Handles document parsing and text extraction with OCR preprocessing."""
    
    def __init__(self, name: str = "OCR module"):
        self.name = name
    
    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file using OCR."""
        try:
            logger.info(f"🔄 Processing PDF with OCR: {file_path}")
            
            # Convert PDF to base64 images
            encoded_images = crop_all_pages(file_path)
            
            if not encoded_images:
                logger.warning("⚠️ No images extracted from PDF, falling back to traditional method")
                return self._extract_text_from_pdf_traditional(file_path)
            
            # Use OCR to extract text from images
            async def extract_text_with_ocr():
                ocr_text = await ocr_list_pages(encoded_images)
                return ocr_text
            
            # Get OCR text
            extracted_text = asyncio.run(extract_text_with_ocr())
            
            if extracted_text and extracted_text.strip():
                logger.info(f"✅ Extracted text from PDF using OCR: {len(extracted_text)} characters")
                return extracted_text.strip()
            else:
                logger.warning("⚠️ OCR extraction failed, falling back to traditional method")
                return self._extract_text_from_pdf_traditional(file_path)
            
        except Exception as e:
            logger.error(f"❌ Error extracting text from PDF with OCR: {e}")
            logger.info("🔄 Falling back to traditional PDF extraction")
            return self._extract_text_from_pdf_traditional(file_path)
    
    def _extract_text_from_pdf_traditional(self, file_path: str) -> str:
        """Traditional PDF text extraction as fallback."""
        try:
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = pypdf.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            
            logger.info(f"✅ Extracted text from PDF (traditional): {len(text)} characters")
            return text.strip()
            
        except Exception as e:
            logger.error(f"❌ Error extracting text from PDF (traditional): {e}")
            return ""
    
    def extract_text_from_image(self, file_path: str) -> str:
        """Extract text from image file using Nanonets OCR."""
        try:
            logger.info(f"🔄 Processing image with OCR: {file_path}")
            encoded_image = encode_image(file_path)
            extracted_text = asyncio.run(ocr_list_pages([encoded_image]))
            logger.info(f"✅ Extracted text from image: {len(extracted_text)} characters")
            return extracted_text
            
        except Exception as e:
            logger.error(f"❌ Error extracting text from image: {e}")
            return ""
    
    def extract_text_from_html(self, file_path: str) -> str:
        """Extract text from HTML file."""
        try:
            return self.nanonets_processor.process_html_file(file_path)
        except Exception as e:
            logger.error(f"❌ Error extracting text from HTML: {e}")
            return ""
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            logger.info(f"✅ Extracted text from DOCX: {len(text)} characters")
            return text.strip()
            
        except Exception as e:
            logger.error(f"❌ Error extracting text from DOCX: {e}")
            return ""
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from TXT file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            logger.info(f"✅ Extracted text from TXT: {len(text)} characters")
            return text.strip()
            
        except Exception as e:
            logger.error(f"❌ Error extracting text from TXT: {e}")
            return ""
    
    def extract_text(self, file_path: str) -> str:
        """Extract text from supported file formats."""
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_ext == '.docx':
            return self.extract_text_from_docx(file_path)
        elif file_ext == '.txt':
            return self.extract_text_from_txt(file_path)
        elif file_ext in ['.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp']:
            return self.extract_text_from_image(file_path)
        elif file_ext in ['.html', '.htm']:
            return self.extract_text_from_html(file_path)
        else:
            logger.error(f"❌ Unsupported file format: {file_ext}")
            return ""

